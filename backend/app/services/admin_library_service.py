"""Admin Documents page: supporting files for the study, two project notes
rendered as text, and the platform's reference data sets. Admin-only — the
routes use require_admin.

- Files (AdminDocument): uploaded by an administrator, encrypted at rest.
  .xlsx / .csv open in the app as a table (one tab per sheet); .pdf and
  images open in a new tab; .docx downloads.
- Notes: only the names in NOTES — a whitelist, so no path from the request
  ever reaches the filesystem (never .env, backups/ or anything else).
- Data sets: categories → items → specs, NCR addresses, market rules,
  withdrawal reasons, document types, profile visibility levels.
"""
import csv
import io
import os
import uuid
from typing import List

from fastapi import HTTPException, UploadFile, status
from sqlmodel import Session, select

from app.core import file_crypto
from app.core.addresses import ncr_addresses_out
from app.core.categories import category_tree_out
from app.core import platform_settings
from app.core.database import engine
from app.core.file_validation import validate_upload
from app.models.admin_document import AdminDocument
from app.models.user import User

BACKEND_ROOT = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
PROJECT_ROOT = os.path.dirname(BACKEND_ROOT)
UPLOAD_DIR = os.path.join("uploads", "admin_documents")

CATEGORIES = {
    "interview_data": "Interview data",
    "thesis": "Thesis chapter",
    "survey": "Survey",
    "other": "Other",
}
ALLOWED_EXTENSIONS = {".xlsx", ".csv", ".pdf", ".png", ".jpg", ".jpeg", ".docx"}
TABLE_EXTENSIONS = {".xlsx", ".csv"}
MAX_BYTES = 25 * 1024 * 1024
# Shown in the app, so a huge sheet doesn't freeze the page.
MAX_TABLE_ROWS = 2000
MAX_TABLE_COLS = 60

# The project notes an admin can read — name → file. Nothing else is servable.
NOTES = {
    "SYSTEM_PROCESS.md": os.path.join(PROJECT_ROOT, "SYSTEM_PROCESS.md"),
    "CONNECT_NOTES.md": os.path.join(PROJECT_ROOT, "CONNECT_NOTES.md"),
}

MEDIA_TYPES = {
    ".pdf": "application/pdf", ".png": "image/png", ".jpg": "image/jpeg", ".jpeg": "image/jpeg",
    ".csv": "text/csv",
    ".xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    ".docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
}


def _not_found() -> HTTPException:
    return HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Not found")


def _ext(filename: str) -> str:
    return os.path.splitext(filename or "")[1].lower()


def _cell(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value)


def parse_table(filename: str, contents: bytes) -> List[dict]:
    """[{name, rows: [[str]], total_rows, truncated}] — one per sheet (a .csv is one sheet)."""
    ext = _ext(filename)
    if ext == ".csv":
        try:
            text = contents.decode("utf-8-sig")
        except UnicodeDecodeError:
            text = contents.decode("latin-1")
        rows = [r[:MAX_TABLE_COLS] for r in csv.reader(io.StringIO(text))]
        return [{"name": os.path.splitext(filename)[0] or "Sheet", "rows": rows[:MAX_TABLE_ROWS],
                 "total_rows": len(rows), "truncated": len(rows) > MAX_TABLE_ROWS}]
    if ext == ".xlsx":
        from openpyxl import load_workbook

        try:
            wb = load_workbook(io.BytesIO(contents), read_only=True, data_only=True)
        except Exception:
            raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="This spreadsheet can't be opened.")
        sheets = []
        for ws in wb.worksheets:
            rows, total = [], 0
            for row in ws.iter_rows(values_only=True):
                cells = [_cell(v) for v in row[:MAX_TABLE_COLS]]
                if not any(cells):
                    continue
                total += 1
                if len(rows) < MAX_TABLE_ROWS:
                    rows.append(cells)
            width = max((len(r) for r in rows), default=0)
            # Trim columns that are empty in every row (a common xlsx artefact).
            while width and all(len(r) < width or not r[width - 1] for r in rows):
                width -= 1
            sheets.append({"name": ws.title, "rows": [r[:width] for r in rows], "total_rows": total,
                           "truncated": total > MAX_TABLE_ROWS})
        wb.close()
        return sheets
    raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="Only .xlsx and .csv files open as a table.")


class AdminLibraryService:
    def _out(self, doc: AdminDocument, uploader: str) -> dict:
        return {
            "id": doc.id, "title": doc.title, "category": doc.category,
            "category_label": CATEGORIES.get(doc.category, doc.category), "filename": doc.filename,
            "size_bytes": doc.size_bytes, "uploaded_by": uploader, "created_at": doc.created_at,
            "opens_as_table": _ext(doc.filename) in TABLE_EXTENSIONS,
        }

    def list(self) -> List[dict]:
        with Session(engine) as session:
            docs = session.exec(select(AdminDocument).order_by(AdminDocument.created_at.desc(), AdminDocument.id.desc())).all()
            names = {u.id: u.business_name for u in session.exec(select(User).where(User.is_admin == True)).all()}  # noqa: E712
            return [self._out(d, names.get(d.uploaded_by, "Administrator")) for d in docs]

    async def upload(self, admin_id: int, title: str, category: str, file: UploadFile) -> dict:
        title = (title or "").strip()
        if len(title) < 2:
            raise HTTPException(status_code=status.HTTP_422_UNPROCESSABLE_ENTITY, detail="Give the document a title.")
        if category not in CATEGORIES:
            raise HTTPException(status_code=status.HTTP_422_UNPROCESSABLE_ENTITY, detail="Pick a category.")
        contents = await file.read()
        ext = validate_upload(file.filename, contents, ALLOWED_EXTENSIONS, MAX_BYTES)
        if ext in TABLE_EXTENSIONS:
            parse_table(file.filename, contents)  # refuse a spreadsheet that won't open

        os.makedirs(UPLOAD_DIR, exist_ok=True)
        path = os.path.join(UPLOAD_DIR, f"{uuid.uuid4().hex}{ext}")
        with open(path, "wb") as f:
            f.write(file_crypto.encrypt_bytes(contents))
        original = os.path.basename(file.filename or f"document{ext}")[:255]
        with Session(engine) as session:
            doc = AdminDocument(title=title[:200], category=category, filename=original, file_path=path,
                                size_bytes=len(contents), uploaded_by=admin_id)
            session.add(doc)
            session.commit()
            session.refresh(doc)
            admin = session.get(User, admin_id)
            return self._out(doc, admin.business_name if admin else "Administrator")

    def _load(self, document_id: int) -> tuple[AdminDocument, bytes]:
        with Session(engine) as session:
            doc = session.get(AdminDocument, document_id)
            if not doc or not os.path.isfile(doc.file_path):
                raise _not_found()
            with open(doc.file_path, "rb") as f:
                return doc, file_crypto.decrypt_bytes_lenient(f.read())

    def file(self, document_id: int) -> tuple[bytes, str, str]:
        doc, contents = self._load(document_id)
        return contents, doc.filename, MEDIA_TYPES.get(_ext(doc.filename), "application/octet-stream")

    def table(self, document_id: int) -> dict:
        doc, contents = self._load(document_id)
        return {"id": doc.id, "title": doc.title, "filename": doc.filename, "sheets": parse_table(doc.filename, contents)}

    def delete(self, document_id: int) -> None:
        with Session(engine) as session:
            doc = session.get(AdminDocument, document_id)
            if not doc:
                raise _not_found()
            path = doc.file_path
            session.delete(doc)
            session.commit()
        if os.path.isfile(path):
            os.remove(path)

    # ---------- notes ----------

    def note_names(self) -> List[str]:
        return [name for name, path in NOTES.items() if os.path.isfile(path)]

    def note(self, name: str) -> dict:
        path = NOTES.get(name)  # whitelist lookup: the request never builds a path
        if not path or not os.path.isfile(path):
            raise _not_found()
        with open(path, encoding="utf-8") as f:
            return {"name": name, "content": f.read()}

    # ---------- data sets ----------

    def datasets(self) -> dict:
        from app.services import fulfillment_service, profile_privacy, withdrawal_policy
        from app.services.business_service import VALID_DOC_TYPES

        doc_labels = {"DTI": "DTI certificate of business name registration (sole proprietorship)",
                      "SEC": "SEC certificate of registration (partnership, corporation, cooperative)",
                      "BIR": "BIR certificate of registration (Form 2303)",
                      "MAYORS_PERMIT": "Mayor's / business permit (optional, toward Tier 3)"}
        return {
            "categories": category_tree_out(),
            "ncr_addresses": ncr_addresses_out(),
            "market_rules": [
                {"label": "Quotations needed to release", "value": str(platform_settings.min_quotations())},
                {"label": "Auto-extension when fewer arrive", "value": f"{platform_settings.auto_extend_days()} days"},
                {"label": "Withdrawal window counted", "value": f"{platform_settings.withdrawal_window_days()} days"},
                {"label": "Order auto-confirms after marked completed",
                 "value": f"{fulfillment_service.AUTO_CONFIRM_DAYS} days"},
                {"label": "Order marked Unconfirmed after no updates",
                 "value": f"{fulfillment_service.UNCONFIRMED_DAYS} days past the delivery date"},
            ],
            "withdrawal_reasons": [{"code": k, "label": v} for k, v in withdrawal_policy.REASONS.items()],
            "withdrawal_consequences": [{"count": n, "stage": s, "text": t}
                                        for n, (s, t) in withdrawal_policy.CONSEQUENCES.items()],
            "document_types": [{"code": c, "label": doc_labels.get(c, c)} for c in sorted(VALID_DOC_TYPES)],
            "profile_visibility": {"sections": list(profile_privacy.SECTIONS), "levels": list(profile_privacy.LEVELS)},
        }


admin_library_service = AdminLibraryService()
